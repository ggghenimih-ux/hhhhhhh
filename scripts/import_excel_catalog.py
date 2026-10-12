#!/usr/bin/env python3
"""Import product rows and embedded images from the September 2026 workbooks.

Dry-run:
  python scripts/import_excel_catalog.py "C:\\path\\to\\catalog.xlsx" ...

Apply:
  python scripts/import_excel_catalog.py --apply "C:\\path\\to\\catalog.xlsx" ...

Requires openpyxl (see scripts/requirements.txt). Incomplete rows are reported
and skipped; the Prisma importer updates matching products and adds new ones.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
IMAGE_DIR = ROOT / "public" / "uploads" / "products"
IMPORTER = ROOT / "prisma" / "importExcelCatalog.ts"
NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "drawing": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

CATEGORIES = [
    ("soldering-hot-air", "Soudure & Air Chaud"),
    ("microscopes", "Microscopes & Optique"),
    ("measuring-power-supply", "Alimentation & Mesure"),
    ("programming-flashing", "Programmation & Flashing"),
    ("screen-refurbish", "Séparateurs & Rénovation Écran"),
    ("equipments-tools", "Outils & Équipements"),
    ("ic-chips", "Puces & Circuits Intégrés (IC)"),
    ("spare-parts", "Pièces Détachées GSM"),
    ("consumables", "Consommables & Chimie"),
    ("accessories", "Accessoires & Câbles"),
    ("other-products", "Autres produits"),
]


def category_for(name: str) -> tuple[str, str]:
    text = name.casefold()
    rules = [
        ("consumables", ("consumables", "solvent", "alcohol", "glue", "flux", "wick", "sponge", "lint-free", "cleaner", "cleaning", "cotton", "擦", "酒精", "无尘布", "洁净棉")),
        ("programming-flashing", ("programmer", "programming", "flashing", "flash tool", "face id", "truetone", "true tone", "activation fixture", "jc id", "jcid", "restore", "编程", "激活")),
        ("screen-refurbish", ("screen", "lcd", "display", "separator", "laminating", "defoaming", "hydrogel", "film", "glass", "screen removal", "贴膜", "分离机", "屏幕")),
        ("microscopes", ("microscope", "microscopy", "trinocular", "eyecup", "auxiliary lens", "magnifier", "optical", "显微镜", "目镜")),
        ("measuring-power-supply", ("power supply", "multimeter", "voltmeter", "ammeter", "meter", "tester", "current", "voltage", "万用表", "电流表", "电源")),
        ("ic-chips", ("bga", "chip", "cpu", "ic ", " ic", "stencil", "steel mesh", "植锡", "钢网", "芯片")),
        ("soldering-hot-air", ("solder", "hot air", "hot-air", "heat gun", "soldering iron", "iron tip", "solder station", "焊台", "热风枪", "烙铁")),
        ("accessories", ("charger", "charging cable", "charging wire", "adapter", "socket", "cable", "connector", "充电", "数据线", "插座")),
        ("spare-parts", ("spare part", "replacement part", "battery", "flex cable", "电池", "排线")),
    ]
    for slug, words in rules:
        if any(word in text for word in words):
            return next((item for item in CATEGORIES if item[0] == slug), CATEGORIES[-1])
    return next(item for item in CATEGORIES if item[0] == "equipments-tools")


def relationship_map(archive: zipfile.ZipFile, source_part: str) -> dict[str, str]:
    rels_part = str(PurePosixPath(source_part).parent / "_rels" / f"{PurePosixPath(source_part).name}.rels")
    if rels_part not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rels_part))
    base = PurePosixPath(source_part).parent
    result: dict[str, str] = {}
    for item in root.findall("pkg:Relationship", NS):
        target = item.attrib.get("Target", "")
        if item.attrib.get("TargetMode") == "External":
            continue
        result[item.attrib["Id"]] = posixpath.normpath(str(base / target))
    return result


def extract_media(archive: zipfile.ZipFile, file_index: int) -> dict[str, str]:
    result: dict[str, str] = {}
    for media_part in archive.namelist():
        if not media_part.startswith("xl/media/") or media_part.endswith("/"):
            continue
        info = archive.getinfo(media_part)
        if info.file_size > 20 * 1024 * 1024:
            raise ValueError(f"Embedded image is unexpectedly large: {media_part}")
        data = archive.read(media_part)
        if data.startswith(b"\xff\xd8\xff"):
            suffix = ".jpeg"
        elif data.startswith(b"\x89PNG\r\n\x1a\n"):
            suffix = ".png"
        elif data.startswith((b"GIF87a", b"GIF89a")):
            suffix = ".gif"
        elif data.startswith(b"BM"):
            suffix = ".bmp"
        elif len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
            suffix = ".webp"
        else:
            raise ValueError(f"Unsupported or invalid embedded raster image: {media_part}")
        digest = hashlib.sha256(data).hexdigest()[:10]
        image_name = f"catalog_{file_index:02d}_{digest}_{Path(media_part).stem}{suffix}"
        result[media_part] = image_name
    return result


def image_anchors(
    archive: zipfile.ZipFile,
    sheet_part: str,
    media_names: dict[str, str],
) -> dict[int, list[str]]:
    anchors: dict[int, list[str]] = {}
    for drawing_part in relationship_map(archive, sheet_part).values():
        if "/drawings/" not in f"/{drawing_part}" or not drawing_part.endswith(".xml"):
            continue
        drawing_rels = relationship_map(archive, drawing_part)
        root = ET.fromstring(archive.read(drawing_part))
        for anchor in list(root):
            start = anchor.find("xdr:from", NS)
            picture = anchor.find(".//drawing:blip", NS)
            if start is None or picture is None:
                continue
            row_text = start.findtext("xdr:row", namespaces=NS)
            media_part = drawing_rels.get(picture.attrib.get(f"{{{NS['rel']}}}embed", ""))
            image_name = media_names.get(media_part or "")
            if row_text is None or image_name is None:
                continue
            anchors.setdefault(int(row_text) + 1, []).append(f"/uploads/products/{image_name}")
    return anchors


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def parse_workbook(
    path: Path,
    file_index: int,
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, bytes], int]:
    products: list[dict[str, object]] = []
    issues: list[dict[str, object]] = []
    images: dict[str, bytes] = {}
    archive = zipfile.ZipFile(path)
    try:
        media_names = extract_media(archive, file_index)
        for media_part, image_name in media_names.items():
            images[image_name] = archive.read(media_part)

        book_relationships = relationship_map(archive, "xl/workbook.xml")
        workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet_parts = {
            sheet.attrib["name"]: book_relationships[sheet.attrib[f"{{{NS['rel']}}}id"]]
            for sheet in workbook_xml.findall(".//main:sheets/main:sheet", NS)
        }
        workbook = load_workbook(path, data_only=True, read_only=False)
        try:
            for sheet in workbook.worksheets:
                numeric_counts = {
                    col: sum(
                        isinstance(sheet.cell(row, col).value, (int, float))
                        and not isinstance(sheet.cell(row, col).value, bool)
                        for row in range(1, sheet.max_row + 1)
                    )
                    for col in range(1, sheet.max_column + 1)
                }
                text_counts = {
                    col: sum(
                        isinstance(sheet.cell(row, col).value, str)
                        and len(sheet.cell(row, col).value.strip()) > 2
                        for row in range(1, sheet.max_row + 1)
                    )
                    for col in range(1, sheet.max_column + 1)
                }
                price_col = max(numeric_counts, key=numeric_counts.get)
                name_col = max(text_counts, key=text_counts.get)
                if numeric_counts[price_col] == 0 or text_counts[name_col] == 0:
                    issues.append({
                        "file": path.name,
                        "sheet": sheet.title,
                        "row": None,
                        "reason": "Could not identify name and price columns",
                    })
                    continue

                price_rows = [
                    row for row in range(1, sheet.max_row + 1)
                    if isinstance(sheet.cell(row, price_col).value, (int, float))
                    and not isinstance(sheet.cell(row, price_col).value, bool)
                    and sheet.cell(row, price_col).value > 0
                ]
                anchors = image_anchors(archive, sheet_parts[sheet.title], media_names)
                anchor_count = sum(len(values) for values in anchors.values())

                for row in price_rows:
                    first_row, last_row = row, row
                    for merged in sheet.merged_cells.ranges:
                        if merged.min_row <= row <= merged.max_row:
                            grouped_prices = [
                                price_row for price_row in price_rows
                                if merged.min_row <= price_row <= merged.max_row
                            ]
                            if len(grouped_prices) == 1:
                                first_row, last_row = merged.min_row, merged.max_row
                                break

                    name = clean_text(" ".join(
                        str(sheet.cell(name_row, name_col).value)
                        for name_row in range(first_row, last_row + 1)
                        if sheet.cell(name_row, name_col).value is not None
                    ))
                    row_images = list(dict.fromkeys(
                        image
                        for image_row in range(first_row, last_row + 1)
                        for image in anchors.get(image_row, [])
                    ))
                    price = float(sheet.cell(row, price_col).value)
                    issue_reasons = []
                    if not name:
                        issue_reasons.append("missing product name")
                    if not row_images:
                        issue_reasons.append("no image anchored to this product row/group")
                    if not price > 0:
                        issue_reasons.append("invalid price")
                    if issue_reasons:
                        issues.append({
                            "file": path.name,
                            "sheet": sheet.title,
                            "row": row,
                            "name": name or None,
                            "price": price,
                            "reason": "; ".join(issue_reasons),
                        })
                        continue

                    category_slug, category_name = category_for(name)
                    products.append({
                        "name": name,
                        "priceTND": price,
                        "shortDescription": "",
                        "description": "",
                        "category": category_slug,
                        "categoryName": category_name,
                        "image": row_images[0],
                        "images": row_images,
                        "stock": 0,
                        "source": f"{path.name}:{sheet.title}:{row}",
                    })
                print(
                    f"  {path.name} / {sheet.title}: {len(price_rows)} priced rows, "
                    f"{anchor_count} embedded image anchors",
                    file=sys.stderr,
                )
        finally:
            workbook.close()
    finally:
        archive.close()
    return products, issues, images, len(media_names)


def deduplicate_products(products: list[dict[str, object]]) -> tuple[list[dict[str, object]], list[str]]:
    selected: dict[str, dict[str, object]] = {}
    duplicates: list[str] = []
    for product in products:
        key = re.sub(r"\s+", " ", str(product["name"]).casefold()).strip()
        if key in selected:
            duplicates.append(str(product["name"]))
        selected[key] = product
    return list(selected.values()), duplicates


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Safely import products from the attached Excel price lists.")
    parser.add_argument("--apply", action="store_true", help="Extract images and import products into the configured Prisma database.")
    parser.add_argument("workbooks", nargs="+", type=Path, help="Excel workbook paths in oldest-to-newest order.")
    args = parser.parse_args()

    for path in args.workbooks:
        if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            parser.error(f"Not an accessible .xlsx/.xlsm workbook: {path}")

    parsed: list[dict[str, object]] = []
    issues: list[dict[str, object]] = []
    media: dict[str, bytes] = {}
    media_count = 0
    for index, path in enumerate(args.workbooks, start=1):
        workbook_products, workbook_issues, workbook_media, count = parse_workbook(path, index)
        parsed.extend(workbook_products)
        issues.extend(workbook_issues)
        media.update(workbook_media)
        media_count += count

    products, duplicates = deduplicate_products(parsed)
    print(json.dumps({
        "workbooks": len(args.workbooks),
        "embeddedImagesExtractable": media_count,
        "imagesUsedByValidProducts": len({image for product in products for image in product["images"]}),
        "validPricedRows": len(parsed),
        "uniqueProducts": len(products),
        "duplicateNamesCollapsed": len(duplicates),
        "incompleteRowsSkipped": len(issues),
        "categories": {
            slug: sum(product["category"] == slug for product in products)
            for slug, _name in CATEGORIES
            if any(product["category"] == slug for product in products)
        },
        "descriptionNote": "The source workbooks contain no product-description column; descriptions are left empty.",
        "issues": issues,
    }, ensure_ascii=False, indent=2))

    if not args.apply:
        print("Dry-run only. No images or database records were changed.")
        return 0
    if not products:
        print("No complete product rows found; refusing to modify files or database.", file=sys.stderr)
        return 1

    created_images: list[Path] = []
    staging_path: Path | None = None
    try:
        IMAGE_DIR.mkdir(parents=True, exist_ok=True)
        for filename, data in media.items():
            target = IMAGE_DIR / filename
            if target.exists():
                if target.read_bytes() != data:
                    raise FileExistsError(f"Existing catalog image differs from source: {target}")
                continue
            target.write_bytes(data)
            created_images.append(target)

        manifest = {
            "products": products,
            "categories": [{"slug": slug, "name": name} for slug, name in CATEGORIES],
        }
        with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", suffix=".json", prefix="catalog-import-", delete=False
        ) as staged:
            json.dump(manifest, staged, ensure_ascii=False)
            staging_path = Path(staged.name)

        npm = shutil.which("npm.cmd") or shutil.which("npm")
        if not npm:
            raise RuntimeError("npm was not found; the database import was not started.")
        command = f'"{npm}" exec -- tsx "{IMPORTER}" --manifest "{staging_path}"'
        result = subprocess.run(command, cwd=ROOT, shell=True, check=False)
        if result.returncode != 0:
            raise RuntimeError(f"Prisma catalog import exited with status {result.returncode}.")
        print(
            f"Import complete. Verified {len(media)} source images "
            f"({len(created_images)} newly extracted) in {IMAGE_DIR}."
        )
        return 0
    except Exception as error:
        for image_path in created_images:
            image_path.unlink(missing_ok=True)
        print(f"Import failed; newly extracted images were rolled back: {error}", file=sys.stderr)
        return 1
    finally:
        if staging_path:
            staging_path.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())
